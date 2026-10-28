"""Geração dos relatórios de conferência, execução e auditoria.

Todos os CSV saem com separador ";" e BOM UTF-8, que é o que o Excel em
português abre com as colunas já separadas e os acentos corretos.
"""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Sequence

from .modelos import Exclusao, ResultadoLeitura, Titulo
from .normalizacao import data_br, formatar_dinheiro_br

_DELIMITADOR = ";"
_CODIFICACAO = "utf-8-sig"


def _escrever(destino: Path, campos: Sequence[str], linhas: Iterable[dict]) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", newline="", encoding=_CODIFICACAO) as arquivo:
        escritor = csv.DictWriter(
            arquivo, fieldnames=list(campos), delimiter=_DELIMITADOR, extrasaction="ignore"
        )
        escritor.writeheader()
        escritor.writerows(linhas)
    return destino


def _detalhe_rateios(titulo: Titulo) -> str:
    return " || ".join(
        f"L{r.linha}: R$ {r.valor_br} ({r.percentual_br}%) "
        f"| CC={r.centro_custo or '[vazio]'} "
        f"| categoria={r.categoria_detalhada or '[vazio]'} "
        f"| subprojeto={r.subprojeto or '[VAZIO - PENDENCIA]'} "
        f"| atividade={r.atividade or '[vazia]'}"
        for r in titulo.rateios
    )


CAMPOS_PRE_EXECUCAO = [
    "situacao", "linhas_planilha", "fornecedor", "documento", "emissao", "vencimento",
    "efetivacao", "competencia", "valor_total", "conta_pagadora", "tipo_despesa",
    "tipo_documento", "tipo_pagamento", "historico", "qtd_rateios", "rateios",
    "duplicidade", "bloqueios", "alertas", "chave", "hash_origem",
]


def relatorio_pre_execucao(resultado: ResultadoLeitura, destino: Path) -> Path:
    """Uma linha por título, com tudo que o usuário precisa conferir (§21)."""
    linhas = []
    for titulo in resultado.titulos:
        datas = titulo.datas_br()
        linhas.append({
            "situacao": "BLOQUEADO" if titulo.bloqueado else "LIBERADO",
            "linhas_planilha": ", ".join(str(x) for x in titulo.linhas_origem),
            "fornecedor": titulo.fornecedor,
            "documento": titulo.documento,
            "emissao": datas["emissao"],
            "vencimento": datas["vencimento"],
            "efetivacao": datas["efetivacao"],
            "competencia": titulo.competencia,
            "valor_total": titulo.valor_br,
            "conta_pagadora": (
                f"{titulo.conta_pagadora_id} - {titulo.conta_pagadora_nome}"
                if titulo.conta_pagadora_id
                else "SEM CONTA PAGADORA"
            ),
            "tipo_despesa": titulo.tipo_despesa,
            "tipo_documento": titulo.tipo_documento,
            "tipo_pagamento": titulo.tipo_pagamento,
            "historico": titulo.historico,
            "qtd_rateios": len(titulo.rateios),
            "rateios": _detalhe_rateios(titulo),
            "duplicidade": titulo.duplicidade,
            "bloqueios": " || ".join(titulo.bloqueios),
            "alertas": " || ".join(titulo.alertas),
            "chave": titulo.chave,
            "hash_origem": titulo.hash_origem(),
        })
    return _escrever(destino, CAMPOS_PRE_EXECUCAO, linhas)


CAMPOS_EXCLUSOES = ["linha_planilha", "fornecedor", "documento", "valor", "motivo"]


def relatorio_exclusoes(exclusoes: list[Exclusao], destino: Path) -> Path:
    """Toda linha descartada com o motivo, exigência do §6."""
    return _escrever(
        destino,
        CAMPOS_EXCLUSOES,
        [
            {
                "linha_planilha": x.linha,
                "fornecedor": x.fornecedor,
                "documento": x.documento,
                "valor": x.valor,
                "motivo": x.motivo,
            }
            for x in sorted(exclusoes, key=lambda e: e.linha)
        ],
    )


CAMPOS_RESULTADO = [
    "situacao", "fornecedor", "documento_original", "documento_final", "id_ongsys",
    "valor", "inicio", "fim", "tentativas", "fornecedor_selecionado",
    "campos_nao_localizados", "avisos", "mensagem", "screenshot",
]


def relatorio_resultado(registros: list[dict], destino: Path) -> Path:
    """Resultado de cada tentativa de lançamento (§21)."""
    return _escrever(destino, CAMPOS_RESULTADO, registros)


CAMPOS_AUDITORIA = [
    "situacao", "id_ongsys", "fornecedor", "documento", "campo", "esperado",
    "encontrado", "confere", "acao_recomendada",
]


def relatorio_auditoria(registros: list[dict], destino: Path) -> Path:
    """Comparação esperado x salvo, campo a campo (§19/§21)."""
    return _escrever(destino, CAMPOS_AUDITORIA, registros)


def resumo_texto(
    resultado: ResultadoLeitura, inicio, fim, conta: str, *, compacto: bool = False
) -> str:
    """Resumo em português mostrado no Prompt de Comando."""
    liberados = [t for t in resultado.titulos if not t.bloqueado]
    bloqueados = [t for t in resultado.titulos if t.bloqueado]
    meses = resultado.meses_com_dados()

    if compacto:
        quantidade_alertas = sum(
            len(t.alertas) for t in resultado.titulos if not t.bloqueado
        )
        partes = [
            "-" * 68,
            "RESUMO DA PLANILHA",
            f"Periodo : {data_br(inicio)} a {data_br(fim)}",
            f"Conta   : {conta}",
            f"Fila    : {len(liberados)} titulo(s) | "
            f"R$ {formatar_dinheiro_br(resultado.total_liberado)}",
        ]
        if bloqueados:
            partes.append(f"Bloqueados: {len(bloqueados)} titulo(s)")
            for titulo in bloqueados:
                partes.append(
                    f"  ! {titulo.fornecedor} / {titulo.documento}: "
                    f"{titulo.bloqueios[0]}"
                )
        if quantidade_alertas:
            partes.append(
                f"Avisos  : {quantidade_alertas} (detalhados no relatorio de conferencia)"
            )
        partes.append("-" * 68)
        return "\n".join(partes)

    partes = [
        "=" * 68,
        "RESUMO DA CONFERENCIA",
        "=" * 68,
        f"Periodo solicitado : {data_br(inicio)} a {data_br(fim)}",
        f"Conta pagadora     : {conta}",
        f"Meses com dados    : {', '.join(meses) if meses else 'NENHUM'}",
        "",
        f"Titulos encontrados: {len(resultado.titulos)}",
        f"  liberados        : {len(liberados)}  ->  R$ {formatar_dinheiro_br(resultado.total_liberado)}",
        f"  bloqueados       : {len(bloqueados)}  ->  R$ "
        f"{formatar_dinheiro_br(resultado.total - resultado.total_liberado)}",
        f"Valor total        : R$ {formatar_dinheiro_br(resultado.total)}",
        "",
        f"Linhas com cor de pendente : {resultado.linhas_por_cor_incluidas}",
        f"Linhas de outras cores     : {resultado.linhas_por_cor_excluidas}",
        f"Linhas descartadas         : {len(resultado.exclusoes)} (veja o relatorio de exclusoes)",
    ]

    meses_pedidos = _meses_do_periodo(inicio, fim)
    ausentes = [m for m in meses_pedidos if m not in meses]
    if ausentes:
        partes += [
            "",
            "ATENCAO: nao ha nenhum lancamento pendente nos meses "
            f"{', '.join(ausentes)} dentro do periodo pedido.",
            "O programa nao inventa linhas. Se voce esperava dados nesses meses,",
            "confirme se a planilha esta atualizada e com as cores corretas.",
        ]

    if bloqueados:
        partes += ["", "TITULOS BLOQUEADOS (nao serao lancados):"]
        for titulo in bloqueados:
            partes.append(f"  - {titulo.fornecedor} / {titulo.documento} / R$ {titulo.valor_br}")
            for bloqueio in titulo.bloqueios:
                partes.append(f"      motivo: {bloqueio}")

    alertados = [t for t in resultado.titulos if t.alertas and not t.bloqueado]
    if alertados:
        partes += ["", "AVISOS (o lancamento continua permitido):"]
        for titulo in alertados:
            for alerta in titulo.alertas:
                partes.append(f"  - {titulo.fornecedor} / {titulo.documento}: {alerta}")

    partes.append("=" * 68)
    return "\n".join(partes)


def _meses_do_periodo(inicio, fim) -> list[str]:
    meses = []
    ano, mes = inicio.year, inicio.month
    while (ano, mes) <= (fim.year, fim.month):
        meses.append(f"{mes:02d}/{ano}")
        mes += 1
        if mes > 12:
            mes, ano = 1, ano + 1
    return meses


def carimbo() -> str:
    """Sufixo de data e hora para não sobrescrever relatórios antigos."""
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def exportar_xlsx(csvs: dict[str, Path], destino: Path) -> Path | None:
    """Junta os CSV gerados em um único .xlsx, quando o openpyxl estiver disponível."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None

    livro = Workbook()
    livro.remove(livro.active)
    cabecalho_fonte = Font(bold=True, color="FFFFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="FF44546A")

    for nome_aba, caminho in csvs.items():
        if not caminho.exists():
            continue
        aba = livro.create_sheet(nome_aba[:31])
        with caminho.open(encoding=_CODIFICACAO, newline="") as arquivo:
            for indice, registro in enumerate(csv.reader(arquivo, delimiter=_DELIMITADOR), start=1):
                aba.append(registro)
                if indice == 1:
                    for celula in aba[1]:
                        celula.font = cabecalho_fonte
                        celula.fill = cabecalho_fundo
                        celula.alignment = Alignment(vertical="center")
        aba.freeze_panes = "A2"
        for coluna in aba.columns:
            largura = max((len(str(c.value or "")) for c in coluna[:200]), default=10)
            aba.column_dimensions[coluna[0].column_letter].width = min(max(largura + 2, 10), 60)

    if not livro.sheetnames:
        return None
    livro.save(destino)
    return destino
