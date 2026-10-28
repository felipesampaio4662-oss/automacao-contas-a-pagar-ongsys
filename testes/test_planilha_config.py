"""Testes de detecção de colunas, cores e validação da configuração."""

from __future__ import annotations

import json
from datetime import date

import apoio
import pytest
from apoio import CABECALHO, PENDENTE, criar_planilha, linha

from ongsys import configuracao as C
from ongsys import planilha as P
from openpyxl import Workbook
from openpyxl.styles import PatternFill


class TestDeteccaoDeColunas:
    """§4: as colunas são achadas pelo nome, não por número fixo."""

    def test_encontra_o_cabecalho_na_linha_3(self, tmp_path):
        destino = criar_planilha(tmp_path / "a.xlsx", [linha("F", 100.0)])
        _, mapa, cabecalho, _ = P.ler_linhas(destino)
        assert cabecalho == 3
        assert mapa["fornecedor"] == 6
        assert mapa["saida"] == 12

    def test_encontra_o_cabecalho_em_outra_linha(self, tmp_path):
        destino = criar_planilha(tmp_path / "b.xlsx", [linha("F", 100.0)], linha_cabecalho=1)
        _, _, cabecalho, _ = P.ler_linhas(destino)
        assert cabecalho == 1

    def test_acompanha_colunas_deslocadas(self, tmp_path):
        """Uma coluna inserida antes não pode quebrar a leitura."""
        livro = Workbook()
        aba = livro.active
        for indice, nome in enumerate(CABECALHO, start=3):  # começa na coluna C
            aba.cell(2, indice, nome)
        dados = linha("FORNECEDOR DESLOCADO", 777.0)
        for indice, valor in enumerate(dados, start=3):
            aba.cell(3, indice, valor)
        aba.cell(3, 8).fill = PatternFill("solid", fgColor=PENDENTE)  # coluna do fornecedor
        destino = tmp_path / "c.xlsx"
        livro.save(destino)

        linhas, mapa, _, _ = P.ler_linhas(destino)
        assert mapa["fornecedor"] == 8
        assert linhas[0].fornecedor == "FORNECEDOR DESLOCADO"
        assert linhas[0].cor == PENDENTE

    def test_nomes_alternativos_de_cabecalho(self, tmp_path):
        livro = Workbook()
        aba = livro.active
        alternativos = ["Codigo de despesa", "Atividade", "Data Emissao", "Data Vencimento",
                        "Data Pagamento", "Beneficiario", "Numero documento", "Centro de custo",
                        "Natureza", "Historico", "Credito", "Debito"]
        for indice, nome in enumerate(alternativos, start=1):
            aba.cell(1, indice, nome)
        for indice, valor in enumerate(linha("F", 100.0), start=1):
            aba.cell(2, indice, valor)
        destino = tmp_path / "d.xlsx"
        livro.save(destino)
        _, mapa, _, _ = P.ler_linhas(destino)
        assert {"fornecedor", "saida", "vencimento", "emissao", "efetivacao"} <= set(mapa)

    def test_cabecalho_sem_colunas_essenciais_da_erro_explicativo(self, tmp_path):
        livro = Workbook()
        aba = livro.active
        aba.cell(1, 1, "Fornecedor")
        aba.cell(1, 2, "Observacao")
        destino = tmp_path / "e.xlsx"
        livro.save(destino)
        with pytest.raises(P.ErroPlanilha) as erro:
            P.ler_linhas(destino)
        assert "faltam colunas" in str(erro.value)

    def test_coluna_manual_sobrescreve_a_deteccao(self, tmp_path):
        destino = criar_planilha(tmp_path / "f.xlsx", [linha("F", 100.0)])
        _, mapa, _, _ = P.ler_linhas(destino, colunas_manuais={"saida": "L"})
        assert mapa["saida"] == 12

    def test_letra_de_coluna_invalida_da_erro(self, tmp_path):
        destino = criar_planilha(tmp_path / "g.xlsx", [linha("F", 100.0)])
        with pytest.raises(P.ErroPlanilha):
            P.ler_linhas(destino, colunas_manuais={"saida": "12"})

    def test_arquivo_inexistente_da_erro_claro(self, tmp_path):
        with pytest.raises(P.ErroPlanilha) as erro:
            P.ler_linhas(tmp_path / "nao_existe.xlsx")
        assert "não foi encontrada" in str(erro.value)

    def test_aba_inexistente_lista_as_disponiveis(self, tmp_path):
        destino = criar_planilha(tmp_path / "h.xlsx", [linha("F", 100.0)])
        with pytest.raises(P.ErroPlanilha) as erro:
            P.ler_linhas(destino, aba="Inexistente")
        assert "Planilha1" in str(erro.value)


class _CorFalsa:
    def __init__(self, tipo, rgb=None, theme=None, indexed=None, tint=0.0):
        self.type, self.rgb, self.theme, self.indexed, self.tint = tipo, rgb, theme, indexed, tint


class _PreenchimentoFalso:
    def __init__(self, tipo, **campos):
        self.patternType = None if tipo is None else "solid"
        self.fgColor = None if tipo is None else _CorFalsa(tipo, **campos)


def _celula_falsa(tipo, **campos):
    """Célula mínima com o formato que o openpyxl entrega para `cor_da_celula`."""
    return type("Celula", (), {"fill": _PreenchimentoFalso(tipo, **campos)})()


class TestCores:
    def test_celula_sem_preenchimento(self, tmp_path):
        destino = criar_planilha(tmp_path / "i.xlsx", [linha("F", 100.0)], cores=[None])
        linhas, _, _, _ = P.ler_linhas(destino)
        assert linhas[0].cor == "SEM_COR"

    def test_cor_de_tema_recebe_rotulo_proprio(self):
        """Cor de tema não pode virar string vazia e passar por 'não pendente' calado.

        Testado direto na função porque o openpyxl regrava a cor de tema como
        RGB ao salvar; o caso real vem de arquivos produzidos pelo Excel.
        """
        assert P.cor_da_celula(_celula_falsa("theme", theme=4, tint=0.4)) == "TEMA:4"

    def test_cor_indexada_recebe_rotulo_proprio(self):
        assert P.cor_da_celula(_celula_falsa("indexed", indexed=13)) == "INDEXADA:13"

    def test_sem_padrao_de_preenchimento(self):
        assert P.cor_da_celula(_celula_falsa(None)) == "SEM_COR"

    def test_cor_rgb_vem_em_maiusculas(self):
        assert P.cor_da_celula(_celula_falsa("rgb", rgb="ffffff99")) == "FFFFFF99"


class TestConfiguracao:
    def _gravar(self, tmp_path, dados) -> "Path":
        destino = tmp_path / "config.json"
        destino.write_text(json.dumps(dados), encoding="utf-8")
        return destino

    def _valida(self) -> dict:
        return {
            "ongsys_url": "https://www.ongsys.com.br/app/index.php/contas-pagar/",
            "contas_pagadoras": {"6": "ITAU"},
            "cores_pendentes": ["FFFFFF99"],
            "tipos_despesa": {"servico": "A", "imposto": "B", "produto": "C", "outros": "D"},
        }

    def test_config_minima_valida(self, tmp_path):
        config = C.carregar(self._gravar(tmp_path, self._valida()))
        assert config["limite_titulos_por_execucao"] == 25

    def test_arquivo_ausente_explica_o_que_fazer(self, tmp_path):
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(tmp_path / "nao_existe.json")
        assert "config.example.json" in str(erro.value)

    def test_json_invalido_indica_a_linha(self, tmp_path):
        destino = tmp_path / "config.json"
        destino.write_text('{"ongsys_url": "https://x",}', encoding="utf-8")
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(destino)
        assert "linha" in str(erro.value)

    def test_campo_obrigatorio_ausente(self, tmp_path):
        dados = self._valida()
        del dados["contas_pagadoras"]
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(self._gravar(tmp_path, dados))
        assert "contas_pagadoras" in str(erro.value)

    def test_cores_pendentes_vazio_e_recusado(self, tmp_path):
        dados = self._valida() | {"cores_pendentes": []}
        with pytest.raises(C.ErroConfiguracao):
            C.carregar(self._gravar(tmp_path, dados))

    def test_cor_com_formato_errado(self, tmp_path):
        dados = self._valida() | {"cores_pendentes": ["amarelo"]}
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(self._gravar(tmp_path, dados))
        assert "ARGB" in str(erro.value)

    def test_mesma_cor_pendente_e_lancada_e_recusada(self, tmp_path):
        dados = self._valida() | {"cores_pendentes": ["FFFFFF99"], "cores_ja_lancadas": ["FFFFFF99"]}
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(self._gravar(tmp_path, dados))
        assert "ao mesmo tempo" in str(erro.value)

    def test_tipo_de_despesa_incompleto(self, tmp_path):
        dados = self._valida()
        del dados["tipos_despesa"]["produto"]
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(self._gravar(tmp_path, dados))
        assert "produto" in str(erro.value)

    def test_conta_padrao_fora_da_lista(self, tmp_path):
        dados = self._valida() | {"conta_pagadora_padrao": "999"}
        with pytest.raises(C.ErroConfiguracao) as erro:
            C.carregar(self._gravar(tmp_path, dados))
        assert "999" in str(erro.value)

    def test_url_precisa_ser_https(self, tmp_path):
        dados = self._valida() | {"ongsys_url": "http://inseguro"}
        with pytest.raises(C.ErroConfiguracao):
            C.carregar(self._gravar(tmp_path, dados))

    def test_semelhanca_fora_do_intervalo(self, tmp_path):
        dados = self._valida() | {"semelhanca_minima_fornecedor": "1.5"}
        with pytest.raises(C.ErroConfiguracao):
            C.carregar(self._gravar(tmp_path, dados))

    def test_campos_desconhecidos_sao_relatados(self, tmp_path):
        dados = self._valida() | {"campo_inventado": 1}
        config = C.carregar(self._gravar(tmp_path, dados))
        assert "campo_inventado" in config["_campos_desconhecidos"]

    def test_config_de_exemplo_do_projeto_e_valida(self):
        """O arquivo distribuído precisa passar na própria validação."""
        config = C.carregar(apoio.RAIZ / "config.example.json")
        assert config["contas_pagadoras"]
