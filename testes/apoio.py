"""Construção de planilhas sintéticas para os testes.

Nenhum teste depende de dados reais da organização, exceto o teste de regressão
que só roda quando a planilha de referência estiver presente na máquina.
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

#: Configuração usada pelos testes que precisam de um config completo. O
#: config.json real não vai para o GitHub; num clone limpo vale o exemplo.
CONFIG_PROJETO = RAIZ / "config.json" if (RAIZ / "config.json").exists() else RAIZ / "config.example.json"

PENDENTE = "FFFFFF99"
JA_LANCADA = "FFFFFFCC"

CABECALHO = [
    "Cod.Orc.", "Rateio", "Emissao", "Vencimento", "Pagamento", "Fornecedor",
    "Documento", "CC", "Titulo categoria", "Categoria detalhada", "Entrada", "Saida",
    "Codigo de barras",
]


def config_de_teste(**ajustes) -> dict:
    """Configuração mínima e válida, já no formato que as regras esperam."""
    from decimal import Decimal

    base = {
        "ongsys_url": "https://www.ongsys.com.br/app/index.php/contas-pagar/",
        "contas_pagadoras": {"6": "BANCO ITAU 12345-6 SEDE", "63": "PREPAGO CARTAO 0001"},
        "conta_pagadora_padrao": "",
        "cores_pendentes": [PENDENTE],
        "cores_ja_lancadas": [JA_LANCADA],
        "ignorar_fornecedores": ["TRANSFERENCIA ENTRE CONTAS", "RECARGA CARTAO PREPAGO",
                                 "APLICACAO CDB", "RESGATE APLICACAO", "RENDIMENTO APLICACAO",
                                 "RECEBIMENTO DE RECURSOS", "TARIFA BANCARIA"],
        "ignorar_categorias": [],
        "ignorar_retencoes": ["IRRF", "PCC"],
        "fornecedores_varejo": ["PAPELMAIS", "HORTIFRUTI DA VILA"],
        "apelidos_fornecedores": {},
        "codigos_subprojeto_por_fornecedor": {},
        "confirmar_duplicidade_por_titulo": True,
        "editar_projetados_automaticamente": True,
        "vencimento_igual_efetivacao": [],
        "tipos_despesa": {
            "servico": "Contratação de Serviço(s)",
            "imposto": "Pagamento de Impostos",
            "produto": "Compra de Produto(s)",
            "outros": "Outros",
        },
        "tipos_documento": {},
        "tipos_pagamento": {},
        "meio_pagamento_padrao": "Internet Banking",
        "limite_titulos_por_execucao": 25,
        "limite_valor_total": Decimal("1000000.00"),
        "alerta_valor_titulo": Decimal("50000.00"),
        "semelhanca_minima_fornecedor": Decimal("0.75"),
        "tolerancia_conciliacao": Decimal("0.01"),
        "timeout_ms": 15000,
        "manter_navegador_visivel": True,
    }
    base.update(ajustes)
    return base


def linha(
    fornecedor: str,
    saida: float | None = 100.0,
    *,
    codigo: str = "2.2.1.05",
    atividade: str = "Despesas ADM",
    emissao: date = date(2026, 6, 1),
    vencimento: date = date(2026, 6, 10),
    pagamento: date = date(2026, 6, 10),
    documento: str = "NF 1",
    cc: str = "GESTAO ADM",
    titulo_categoria: str = "DESPESAS_ADMINISTRATIVAS",
    categoria: str = "MATERIAL DE ESCRITORIO",
    entrada: float | None = None,
    codigo_barras: str = "",
) -> list:
    return [codigo, atividade, emissao, vencimento, pagamento, fornecedor,
            documento, cc, titulo_categoria, categoria, entrada, saida, codigo_barras]


def criar_planilha(destino: Path, linhas: list[list], cores: list[str] | None = None,
                   linha_cabecalho: int = 3) -> Path:
    """Grava um .xlsx com o cabeçalho na linha indicada e as cores pedidas."""
    livro = Workbook()
    aba = livro.active
    aba.title = "Planilha1"

    for indice, nome in enumerate(CABECALHO, start=1):
        aba.cell(linha_cabecalho, indice, nome)

    cores = cores or [PENDENTE] * len(linhas)
    for deslocamento, (dados, cor) in enumerate(zip(linhas, cores)):
        numero = linha_cabecalho + 1 + deslocamento
        for coluna, valor in enumerate(dados, start=1):
            aba.cell(numero, coluna, valor)
        if cor:
            aba.cell(numero, 6).fill = PatternFill("solid", fgColor=cor)

    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


def montar(destino: Path, linhas: list[list], cores: list[str] | None = None,
           config: dict | None = None, conta: str = "6",
           inicio: date = date(2026, 1, 1), fim: date = date(2026, 12, 31),
           linha_cabecalho: int = 3):
    """Atalho: cria a planilha, lê e aplica as regras."""
    from ongsys import planilha as P
    from ongsys import regras as R

    criar_planilha(destino, linhas, cores, linha_cabecalho)
    lidas, _, _, _ = P.ler_linhas(destino)
    return R.montar_titulos(lidas, inicio, fim, config or config_de_teste(), conta)
